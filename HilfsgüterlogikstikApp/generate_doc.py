"""Generate project documentation as a Word (.docx) file."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# -- Styles --
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(4)

for level in range(1, 4):
    hs = doc.styles[f"Heading {level}"]
    hs.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

# ============================================================
# TITLE
# ============================================================
title = doc.add_heading("HilfsgüterlogistikApp – Projektdokumentation", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph(
    "Komponentenübersicht, Abhängigkeiten und Schnittstellendokumentation",
    style="Subtitle",
).alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph("")  # spacer

# ============================================================
# PART 1 – KOMPONENTENÜBERSICHT
# ============================================================
doc.add_heading("Teil 1: Komponentenübersicht & Abhängigkeiten", level=1)

# -- Projektstruktur --
doc.add_heading("Projektstruktur", level=2)
p = doc.add_paragraph()
p.add_run("Typ: ").bold = True
p.add_run("Django 6.0 Webanwendung mit Admin-UI (SimpleUI)")
p = doc.add_paragraph()
p.add_run("Datenbank: ").bold = True
p.add_run("SQLite (lokal) / MySQL (Docker)")
p = doc.add_paragraph()
p.add_run("Branch: ").bold = True
p.add_run("prototype")

# -- Django Apps --
doc.add_heading("Django Apps (Komponenten)", level=2)

# App 1
doc.add_heading("HilfsgüterlogikstikApp (Hauptprojekt)", level=3)
doc.add_paragraph(
    "Settings, URL-Routing, WSGI/ASGI-Konfiguration. "
    "Leitet / → /admin/ weiter. Konfiguriert alle anderen Apps."
)

# App 2
doc.add_heading("core – Stammdaten", level=3)
doc.add_paragraph("Status (typ: Order | Inspection | Position)")
doc.add_paragraph(
    "Hinweis: Diese App ist nicht in INSTALLED_APPS registriert. "
    "Das Status-Model wird von keinem anderen Model referenziert.",
    style="List Bullet",
)

# App 3
doc.add_heading("auftraege – Auftragsverwaltung", level=3)
items = [
    ("Order", "Hat einen Status (TextChoices: OFFEN → IN_BEARBEITUNG → IN_PRUEFUNG → GEPRUEFT → ABGESCHLOSSEN → STORNIERT)"),
    ("Container", "Gehört zu einem Order (FK, CASCADE)"),
    ("Box", "Gehört zu einem Container (FK, CASCADE)"),
    ("Item", "Gehört zu einer Box (FK, CASCADE)"),
]
for name, desc in items:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(name).bold = True
    p.add_run(f" – {desc}")

# App 4
doc.add_heading("pruefung – Prüfungswesen", level=3)
items = [
    ("OrderInspection", "Gehört zu einem Order (FK), hat einen User als Inspector"),
    ("IndividualInspection", "Gehört zu einer OrderInspection (FK), referenziert Container (FK), hat User als Inspector"),
    ("InspectionResult", "Gehört zu einer IndividualInspection (FK), referenziert Item (FK)"),
    ("Shrinkage", "1:1-Beziehung zu OrderInspection, hat User als Created By"),
]
for name, desc in items:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(name).bold = True
    p.add_run(f" – {desc}")

# App 5
doc.add_heading("permissions – Rollensystem", level=3)
doc.add_paragraph(
    "Hinweis: Nicht in INSTALLED_APPS registriert – Tabellen existieren nicht in der DB.",
    style="List Bullet",
)
items = [
    ("Rolle", "Rollen im System"),
    ("Permission", "Berechtigungen"),
    ("RolePermission", "Rolle ↔ Permission (n:m über Zwischentabelle)"),
    ("UserRolle", "Django User ↔ Rolle (n:m über Zwischentabelle)"),
]
for name, desc in items:
    p = doc.add_paragraph(style="List Bullet 2" if "List Bullet 2" in [s.name for s in doc.styles] else "List Bullet")
    p.add_run(name).bold = True
    p.add_run(f" – {desc}")

# -- Abhängigkeitsgraph --
doc.add_heading("Abhängigkeitsgraph", level=2)

graph_text = """\
HilfsgüterlogistikApp (Hauptprojekt)
│
├── django.contrib.auth
│   └── User
│
├── core
│   └── Status
│
├── auftraege
│   └── Order
│       └── Container
│           └── Box
│               └── Item
│
├── pruefung  (hängt ab von: auftraege, auth)
│   ├── OrderInspection ──FK──► Order
│   │       └──FK──► User
│   ├── IndividualInspection
│   │   ├──FK──► OrderInspection
│   │   ├──FK──► Container
│   │   └──FK──► User
│   ├── InspectionResult
│   │   ├──FK──► IndividualInspection
│   │   └──FK──► Item
│   └── Shrinkage
│       ├──1:1──► OrderInspection
│       └──FK──► User
│
└── permissions  (hängt ab von: auth) [NICHT REGISTRIERT]
    ├── Rolle
    ├── Permission
    ├── RolePermission ──FK──► Rolle + Permission
    └── UserRolle ──FK──► User + Rolle"""

p = doc.add_paragraph()
run = p.add_run(graph_text)
run.font.name = "Consolas"
run.font.size = Pt(9)

# -- Geschäftslogik --
doc.add_heading("Geschäftslogik-Verbindung", level=2)
doc.add_paragraph(
    "Wenn eine OrderInspection gespeichert wird, wird automatisch "
    "Order.update_status_from_inspection() aufgerufen. Der Order-Status ändert sich "
    "basierend auf dem overall_status der letzten Prüfung."
)

# -- Infrastruktur --
doc.add_heading("Infrastruktur", level=2)
items = [
    "Datenbank: SQLite (lokal) oder MySQL (via Docker, gesteuert durch DB_ENGINE)",
    "Admin-UI: Django Admin mit SimpleUI Theme",
    "Static Files: staticfiles/ mit SimpleUI-Assets, FontAwesome, ElementUI",
    "Templates: Custom base_site.html für Admin-Branding",
]
for item in items:
    doc.add_paragraph(item, style="List Bullet")

# ============================================================
# PART 2 – SCHNITTSTELLENDOKUMENTATION
# ============================================================
doc.add_page_break()
doc.add_heading("Teil 2: Schnittstellendokumentation", level=1)

# -- 1. HTTP --
doc.add_heading("1. HTTP-Schnittstellen (URL-Endpunkte)", level=2)
doc.add_paragraph(
    "Das Projekt hat keine eigenen API-Endpunkte oder Views. "
    "Es gibt nur zwei URLs:"
)

table = doc.add_table(rows=3, cols=3)
table.style = "Light Grid Accent 1"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = table.rows[0].cells
hdr[0].text = "Pfad"
hdr[1].text = "Ziel"
hdr[2].text = "Typ"
row1 = table.rows[1].cells
row1[0].text = "/"
row1[1].text = "RedirectView → /admin/"
row1[2].text = "Permanente Weiterleitung"
row2 = table.rows[2].cells
row2[0].text = "/admin/*"
row2[1].text = "Django Admin Site"
row2[2].text = "Admin CRUD-Interface"

doc.add_paragraph("")
p = doc.add_paragraph()
p.add_run("Server-Protokolle:").bold = True
doc.add_paragraph("WSGI (wsgi.py) – synchrones HTTP (Standard-Deployment)", style="List Bullet")
doc.add_paragraph("ASGI (asgi.py) – asynchrones HTTP (optional)", style="List Bullet")

# -- 2. Datenbank --
doc.add_heading("2. Datenbank-Schnittstellen (Model-Relationen)", level=2)

doc.add_heading("App auftraege – interne Hierarchie", level=3)
table = doc.add_table(rows=4, cols=4)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
hdr[0].text = "Von"
hdr[1].text = "Beziehung"
hdr[2].text = "Nach"
hdr[3].text = "on_delete"
data = [
    ("Order", "1:n", "Container", "CASCADE"),
    ("Container", "1:n", "Box", "CASCADE"),
    ("Box", "1:n", "Item", "CASCADE"),
]
for i, (von, bez, nach, delete) in enumerate(data, 1):
    row = table.rows[i].cells
    row[0].text = von
    row[1].text = bez
    row[2].text = nach
    row[3].text = delete

doc.add_paragraph(
    "Alle Löschungen kaskadieren: Wird ein Order gelöscht, werden alle "
    "zugehörigen Container, Boxes und Items mitgelöscht."
)

doc.add_heading("App pruefung → App auftraege (App-übergreifend)", level=3)
table = doc.add_table(rows=7, cols=5)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
hdr[0].text = "Von"
hdr[1].text = "Beziehung"
hdr[2].text = "Nach"
hdr[3].text = "on_delete"
hdr[4].text = "Nullable"
data = [
    ("OrderInspection", "n:1", "Order", "CASCADE", "Nein"),
    ("IndividualInspection", "n:1", "OrderInspection", "CASCADE", "Nein"),
    ("IndividualInspection", "n:1", "Container", "CASCADE", "Ja"),
    ("InspectionResult", "n:1", "IndividualInspection", "CASCADE", "Nein"),
    ("InspectionResult", "n:1", "Item", "CASCADE", "Nein"),
    ("Shrinkage", "1:1", "OrderInspection", "CASCADE", "Nein"),
]
for i, (von, bez, nach, delete, null) in enumerate(data, 1):
    row = table.rows[i].cells
    row[0].text = von
    row[1].text = bez
    row[2].text = nach
    row[3].text = delete
    row[4].text = null

doc.add_heading("App pruefung → django.contrib.auth", level=3)
table = doc.add_table(rows=4, cols=5)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
hdr[0].text = "Von"
hdr[1].text = "Feld"
hdr[2].text = "Nach"
hdr[3].text = "on_delete"
hdr[4].text = "Nullable"
data = [
    ("OrderInspection", "inspector", "User", "SET_NULL", "Ja"),
    ("IndividualInspection", "inspector", "User", "SET_NULL", "Ja"),
    ("Shrinkage", "created_by", "User", "SET_NULL", "Ja"),
]
for i, (von, feld, nach, delete, null) in enumerate(data, 1):
    row = table.rows[i].cells
    row[0].text = von
    row[1].text = feld
    row[2].text = nach
    row[3].text = delete
    row[4].text = null

doc.add_heading("App permissions → django.contrib.auth", level=3)
doc.add_paragraph("Nicht registriert – existiert nur im Code, nicht in der DB.", style="List Bullet")
table = doc.add_table(rows=5, cols=4)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
hdr[0].text = "Von"
hdr[1].text = "Beziehung"
hdr[2].text = "Nach"
hdr[3].text = "on_delete"
data = [
    ("UserRolle", "n:1", "User", "CASCADE"),
    ("UserRolle", "n:1", "Rolle", "CASCADE"),
    ("RolePermission", "n:1", "Rolle", "CASCADE"),
    ("RolePermission", "n:1", "Permission", "CASCADE"),
]
for i, (von, bez, nach, delete) in enumerate(data, 1):
    row = table.rows[i].cells
    row[0].text = von
    row[1].text = bez
    row[2].text = nach
    row[3].text = delete

# -- 3. Geschäftslogik --
doc.add_heading("3. Geschäftslogik-Schnittstellen (Hooks & Callbacks)", level=2)

doc.add_heading("OrderInspection.save() → Order.update_status_from_inspection()", level=3)
doc.add_paragraph(
    "Beim Speichern einer OrderInspection wird automatisch der Status des "
    "zugehörigen Orders aktualisiert. Dies ist die einzige automatisierte "
    "Geschäftslogik im Projekt."
)
doc.add_paragraph("Quellcode: pruefung/models.py (Zeilen 31–34), auftraege/models.py (Zeilen 30–44)")

doc.add_paragraph("Status-Mapping:")
table = doc.add_table(rows=5, cols=2)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
hdr[0].text = "OrderInspection.overall_status"
hdr[1].text = "→ Order.status"
data = [
    ("OFFEN", "IN_PRUEFUNG"),
    ("IN_PRUEFUNG", "IN_PRUEFUNG"),
    ("BESTANDEN", "GEPRUEFT"),
    ("NICHT_BESTANDEN", "IN_BEARBEITUNG"),
]
for i, (von, nach) in enumerate(data, 1):
    row = table.rows[i].cells
    row[0].text = von
    row[1].text = nach

doc.add_paragraph("Order.status ist im Admin readonly und wird ausschließlich durch diesen Mechanismus gesteuert.")

# -- 4. Admin --
doc.add_heading("4. Admin-Schnittstellen (UI-Logik & dynamische Filter)", level=2)

doc.add_heading("Inline-Verschachtelungen", level=3)
inlines = [
    "OrderAdmin → ContainerInline (TabularInline)",
    "ContainerAdmin → BoxInline (TabularInline, neue Boxes) + ContainerForm (bestehende Boxes zuweisen)",
    "BoxAdmin → ItemInline (TabularInline)",
    "OrderInspectionAdmin → IndividualInspectionInline (TabularInline) + ShrinkageInline (StackedInline, 1:1)",
    "IndividualInspectionAdmin → InspectionResultInline (TabularInline)",
]
for item in inlines:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("Dynamische Queryset-Filter im Admin", level=3)
filters = [
    (
        "IndividualInspectionInline → Container-Feld",
        "pruefung/admin.py (Zeilen 45–52)",
        "Zeigt nur Container des zugehörigen Orders an.",
    ),
    (
        "IndividualInspectionAdmin → Container-Feld",
        "pruefung/admin.py (Zeilen 86–100)",
        "Beim Bearbeiten: Filter auf Container des zugehörigen Orders. "
        "Beim Erstellen via Inline: Filter über request._order_inspection_obj. "
        "Fallback: alle Container.",
    ),
    (
        "InspectionResultInline → Item-Feld",
        "pruefung/admin.py (Zeilen 17–30)",
        "Wenn Container vorhanden: nur Items aus Boxes dieses Containers. "
        "Fallback: alle Items des Orders. Sonst: leeres Queryset.",
    ),
    (
        "ContainerForm → Boxes-Feld",
        "auftraege/admin.py (Zeilen 7–32)",
        "FilteredSelectMultiple-Widget zum Zuweisen bestehender Boxes zu einem Container.",
    ),
]
for title, source, desc in filters:
    p = doc.add_paragraph()
    p.add_run(f"{title}").bold = True
    p.add_run(f" ({source})")
    doc.add_paragraph(desc)

doc.add_heading("Request-basierter Datenaustausch", level=3)
doc.add_paragraph(
    "IndividualInspectionInline.get_formset() setzt request._order_inspection_obj, "
    "welches in formfield_for_foreignkey() für den Container-Filter gelesen wird."
)
doc.add_paragraph(
    "InspectionResultInline.get_formset() setzt request._individual_inspection_obj, "
    "welches in formfield_for_foreignkey() für den Item-Filter gelesen wird."
)

# -- 5. Konfiguration --
doc.add_heading("5. Konfigurations-Schnittstellen (Umgebungsvariablen)", level=2)

table = doc.add_table(rows=12, cols=3)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
hdr[0].text = "Variable"
hdr[1].text = "Beschreibung"
hdr[2].text = "Default"
data = [
    ("DOCKER_ENV", "Erkennung Docker-Umgebung", "(nicht gesetzt)"),
    ("DJANGO_ENV", "Umgebungsname für .env-Datei", "local"),
    ("SECRET_KEY", "Django Secret Key", "(insecure fallback)"),
    ("DEBUG", "Debug-Modus", "0 (aus)"),
    ("ALLOWED_HOSTS", "Komma-getrennte Hostliste", '""'),
    ("DB_ENGINE", "Datenbank-Engine", "(nicht gesetzt → SQLite)"),
    ("DB_NAME", "Datenbankname", "django_db"),
    ("DB_USER", "DB-Benutzer", "django_user"),
    ("DB_PASSWORD", "DB-Passwort", "django_password"),
    ("DB_HOST", "DB-Host", "db"),
    ("DB_PORT", "DB-Port", "3306"),
]
for i, (var, desc, default) in enumerate(data, 1):
    row = table.rows[i].cells
    row[0].text = var
    row[1].text = desc
    row[2].text = default

doc.add_paragraph(
    "Logik: Wenn DOCKER_ENV nicht gesetzt ist, wird .env.{DJANGO_ENV} aus dem "
    "übergeordneten Verzeichnis geladen. Wenn DB_ENGINE gesetzt ist, wird MySQL "
    "verwendet; andernfalls SQLite."
)

# -- 6. Management Commands --
doc.add_heading("6. Management-Command-Schnittstelle", level=2)
p = doc.add_paragraph()
run = p.add_run("python manage.py populate_testdata")
run.bold = True
run.font.name = "Consolas"
doc.add_paragraph("Quelle: auftraege/management/commands/populate_testdata.py")
doc.add_paragraph("Funktionen:", style="List Bullet")
doc.add_paragraph("Löscht alle bestehenden Testdaten (mit SET FOREIGN_KEY_CHECKS=0 – nur MySQL-kompatibel)", style="List Bullet")
doc.add_paragraph("Erstellt: 4 Groups, 4 Users, 4 Orders mit Containern/Boxes/Items, OrderInspections, IndividualInspections, InspectionResults, 1 Shrinkage", style="List Bullet")

# -- 7. Static Assets --
doc.add_heading("7. Statische Assets / Template-Schnittstelle", level=2)
doc.add_paragraph("templates/admin/base_site.html – Überschreibt das Django-Admin-Branding", style="List Bullet")
doc.add_paragraph("static/admin/css/logistik-theme.css – Custom CSS für Admin", style="List Bullet")
doc.add_paragraph("SimpleUI als Admin-Theme (via INSTALLED_APPS)", style="List Bullet")

# ============================================================
# SAVE
# ============================================================
output_path = "/Users/nilsgronowski/workspace/OPS 2026/HilfsgüterlogikstikApp/Projektdokumentation_HilfsgüterlogistikApp.docx"
doc.save(output_path)
print(f"✅ Dokument gespeichert: {output_path}")
